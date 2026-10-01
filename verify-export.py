import importlib.util
import json
import openpyxl
import zipfile
import xml.etree.ElementTree as ET

print({x: bool(importlib.util.find_spec(x)) for x in ['formulas', 'pycel', 'xlsx_evaluate']})
path = 'verification/expanded.xlsx'
with zipfile.ZipFile(path) as z:
    assert z.testzip() is None
    for name in z.namelist():
        ET.fromstring(z.read(name))
w = openpyxl.load_workbook(path, data_only=False)
v = openpyxl.load_workbook(path, data_only=True)
assert w.sheetnames == ['Модель', 'Вводные']
with open('verification/expected.json', encoding='utf-8') as f:
    expected = json.load(f)
count = 0
for row in expected['rows']:
    for i, value in enumerate(row['values']):
        cell = w.worksheets[0].cell(row['excelRow'], i + 2)
        assert cell.value == '=' + row['formulas'][i]
        assert abs(v.worksheets[0].cell(row['excelRow'], i + 2).value - value) < .001
        count += 1
for i, item in enumerate(expected['inputs']):
    assert v.worksheets[1].cell(i + 4, 2).value == item['value']
assert w.worksheets[1].data_validations.count > 0
download = openpyxl.load_workbook('verification/v2-browser.xlsx', data_only=True)
assert download.sheetnames == ['Модель', 'Вводные']
print('Verified formula/cache cells:', count, '; inputs:', len(expected['inputs']))
